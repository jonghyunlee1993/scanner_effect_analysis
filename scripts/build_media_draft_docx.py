#!/usr/bin/env python3
"""Build the Korean Medical Image Analysis working draft as a Word document."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = REPO_ROOT / "docs/manuscript/scanner_spectrum_media_draft_ko.md"
DEFAULT_OUTPUT = REPO_ROOT / "docs/manuscript/scanner_spectrum_media_draft_ko.docx"
FIGURE_RE = re.compile(r"^\[\[FIGURE:(.+?)\|(.+)\]\]$")
INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*)")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def set_keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def set_east_asia_font(style, font_name: str) -> None:
    style.font.name = font_name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)

    normal = document.styles["Normal"]
    set_east_asia_font(normal, "Malgun Gothic")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.25

    style_specs = {
        "Title": (18, RGBColor(31, 55, 80), True),
        "Subtitle": (11, RGBColor(89, 89, 89), False),
        "Heading 1": (15, RGBColor(31, 78, 121), True),
        "Heading 2": (13, RGBColor(47, 84, 150), True),
        "Heading 3": (11.5, RGBColor(68, 68, 68), True),
        "Heading 4": (10.5, RGBColor(68, 68, 68), True),
        "Caption": (9, RGBColor(89, 89, 89), False),
        "List Bullet": (10.5, RGBColor(0, 0, 0), False),
        "List Number": (10.5, RGBColor(0, 0, 0), False),
    }
    for name, (size, color, bold) in style_specs.items():
        style = document.styles[name]
        set_east_asia_font(style, "Malgun Gothic")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = bold

    document.styles["Title"].paragraph_format.space_after = Pt(12)
    for heading_name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        style = document.styles[heading_name]
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    if "Draft Note" not in document.styles:
        note = document.styles.add_style("Draft Note", WD_STYLE_TYPE.PARAGRAPH)
    else:
        note = document.styles["Draft Note"]
    set_east_asia_font(note, "Malgun Gothic")
    note.font.size = Pt(9.5)
    note.font.color.rgb = RGBColor(127, 96, 0)
    note.paragraph_format.left_indent = Cm(0.4)
    note.paragraph_format.right_indent = Cm(0.4)
    note.paragraph_format.space_before = Pt(5)
    note.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.text = "Scanner Spectrum | Medical Image Analysis 한글 working draft"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(128, 128, 128)
    add_page_number(section.footer.paragraphs[0])

    properties = document.core_properties
    properties.title = "Scanner Spectrum — Medical Image Analysis Korean working draft"
    properties.subject = "Frequency-resolved scanner effects in pathology foundation models"
    properties.keywords = "digital pathology; scanner; foundation model; harmonization"
    properties.comments = "Internal Korean working draft generated from the project protocol."


def add_inline(paragraph, text: str, font_size: float | None = None) -> None:
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            if font_size:
                run.font.size = Pt(font_size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
            run.font.color.rgb = RGBColor(80, 80, 80)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        if font_size:
            run.font.size = Pt(font_size)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        if font_size:
            run.font.size = Pt(font_size)


def add_paragraph(document: Document, text: str, style: str | None = None, *, cover=False):
    paragraph = document.add_paragraph(style=style)
    add_inline(paragraph, text)
    if cover:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if style is None:
            paragraph.style = document.styles["Subtitle"]
    return paragraph


def parse_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_table_separator(line: str) -> bool:
    cells = parse_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def add_table(document: Document, rows: list[list[str]]) -> None:
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        if row_index == 0:
            set_repeat_table_header(row)
        for column_index in range(column_count):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            paragraph = cell.paragraphs[0]
            add_inline(paragraph, values[column_index] if column_index < len(values) else "", 8.5)
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.space_before = Pt(1)
            if row_index == 0:
                set_cell_shading(cell, "D9EAF7")
                for run in paragraph.runs:
                    run.bold = True
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F7F9FB")
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_figure(document: Document, figure_path: Path, caption: str) -> None:
    document.add_page_break()
    if not figure_path.exists():
        note = document.add_paragraph(style="Draft Note")
        note.add_run(f"[그림 파일을 찾을 수 없음: {figure_path}]")
        set_cell_like_paragraph_shading(note, "FFF2CC")
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(figure_path), width=Inches(6.55))
    caption_paragraph = document.add_paragraph(style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(caption_paragraph, caption)
    caption_paragraph.paragraph_format.keep_with_next = False


def set_cell_like_paragraph_shading(paragraph, fill: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def render_markdown(document: Document, source: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    on_cover = True
    first_title = True
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        if stripped == "<!-- PAGEBREAK -->":
            document.add_page_break()
            on_cover = False
            index += 1
            continue

        figure_match = FIGURE_RE.match(stripped)
        if figure_match:
            raw_path, caption = figure_match.groups()
            path = Path(raw_path)
            if not path.is_absolute():
                path = REPO_ROOT / path
            add_figure(document, path, caption)
            index += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            heading = stripped[level:].strip()
            if level == 1 and first_title:
                paragraph = add_paragraph(document, heading, "Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                first_title = False
            else:
                style = f"Heading {min(level - 1, 4)}"
                paragraph = add_paragraph(document, heading, style)
                set_keep_with_next(paragraph)
            index += 1
            continue

        if stripped.startswith("> "):
            paragraph = add_paragraph(document, stripped[2:], "Draft Note")
            set_cell_like_paragraph_shading(paragraph, "FFF2CC")
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
            rows = [parse_table_row(stripped)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(parse_table_row(lines[index]))
                index += 1
            add_table(document, rows)
            continue

        bullet_match = re.match(r"^-\s+(.+)$", stripped)
        number_match = re.match(r"^\d+\.\s+(.+)$", stripped)
        if bullet_match:
            add_paragraph(document, bullet_match.group(1), "List Bullet")
            index += 1
            continue
        if number_match:
            add_paragraph(document, number_match.group(1), "List Number")
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate:
                break
            if (
                candidate.startswith("#")
                or candidate.startswith("> ")
                or candidate.startswith("|")
                or candidate.startswith("[[FIGURE:")
                or candidate == "<!-- PAGEBREAK -->"
                or re.match(r"^-\s+", candidate)
                or re.match(r"^\d+\.\s+", candidate)
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        text = " ".join(paragraph_lines)
        paragraph = add_paragraph(document, text, cover=on_cover)
        if text.startswith("T_s(") or text.startswith("y_tsr") or text.startswith("outcome ~") or text.startswith("+"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.name = "Cambria Math"


def add_end_marker(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(12)
    run = paragraph.add_run("— 한글 working draft 끝 —")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(150, 150, 150)


def main() -> None:
    args = parse_args()
    source = args.source.resolve()
    output = args.output.resolve()
    if not source.exists():
        raise FileNotFoundError(source)
    output.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    configure_document(document)
    render_markdown(document, source)
    add_end_marker(document)
    document.save(output)
    print(output)


if __name__ == "__main__":
    main()
